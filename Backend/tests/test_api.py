import io
import json

import openpyxl
import pytest

from app.services import llm
from tests.conftest import SAMPLE_TXT, make_pdf


def test_health_and_docs(client):
    r = client.get("/health")
    assert r.status_code == 200 and r.json()["status"] == "ok"
    assert r.json()["groq_configured"] is True
    assert client.get("/docs").status_code == 200
    assert client.get("/openapi.json").status_code == 200


def test_cors_allows_vite_origins(client):
    for origin in ("http://localhost:5173", "http://127.0.0.1:5173"):
        r = client.options(
            "/api/documents",
            headers={"Origin": origin, "Access-Control-Request-Method": "GET"},
        )
        assert r.headers.get("access-control-allow-origin") == origin


# ---------------------------------------------------------------- documents
def test_upload_txt_and_list(client):
    r = client.post("/api/documents/upload", files=[("files", ("a.txt", SAMPLE_TXT.encode(), "text/plain"))])
    assert r.status_code == 201
    doc = r.json()["documents"][0]
    assert doc["status"] == "ready" and doc["chunk_count"] > 1 and doc["filename"] == "a.txt"
    listing = client.get("/api/documents").json()
    assert listing["total"] == 1 and listing["documents"][0]["id"] == doc["id"]


def test_upload_multiple_files_mixed_valid_invalid(client):
    files = [
        ("files", ("ok.txt", SAMPLE_TXT.encode(), "text/plain")),
        ("files", ("bad.exe", b"MZ", "application/octet-stream")),
        ("files", ("empty.txt", b"", "text/plain")),
    ]
    r = client.post("/api/documents/upload", files=files)
    assert r.status_code == 201
    body = r.json()
    assert body["uploaded"] == 1
    assert {e["code"] for e in body["errors"]} == {"unsupported_type", "empty_file"}


def test_upload_all_invalid_returns_422(client):
    r = client.post("/api/documents/upload", files=[("files", ("x.docx", b"data", "application/octet-stream"))])
    assert r.status_code == 422 and r.json()["code"] == "no_documents_processed"


def test_upload_without_files(client):
    r = client.post("/api/documents/upload")
    assert r.status_code in (400, 422)


def test_upload_pdf_keeps_page_numbers(client, mock_llm):
    pdf = make_pdf(["Alpha page about apples and orchards", "Beta page about vector databases and Chroma"])
    r = client.post("/api/documents/upload", files=[("files", ("paper.pdf", pdf, "application/pdf"))])
    assert r.status_code == 201, r.text
    assert r.json()["documents"][0]["page_count"] == 2
    ans = client.post("/api/ask", json={"question": "Which page covers vector databases and Chroma?"}).json()
    assert ans["sources"][0]["page"] == 2


def test_upload_corrupt_pdf_is_reported(client):
    r = client.post("/api/documents/upload", files=[("files", ("bad.pdf", b"%PDF-1.4 junk", "application/pdf"))])
    assert r.status_code == 422
    docs = client.get("/api/documents").json()["documents"]
    assert docs[0]["status"] == "failed" and docs[0]["error"]


def test_delete_document_removes_chunks(client, uploaded):
    from sqlalchemy import func, select

    from app.database import SessionLocal
    from app.models import Chunk

    with SessionLocal() as db:
        assert db.execute(select(func.count(Chunk.id))).scalar_one() > 0
    client.delete(f"/api/documents/{uploaded['id']}")
    with SessionLocal() as db:
        assert db.execute(select(func.count(Chunk.id))).scalar_one() == 0  # no orphans


def test_delete_document(client, uploaded, mock_llm):
    assert client.delete(f"/api/documents/{uploaded['id']}").status_code == 200
    assert client.get("/api/documents").json()["total"] == 0
    assert client.delete(f"/api/documents/{uploaded['id']}").status_code == 404
    r = client.post("/api/ask", json={"question": "anything?"})
    assert r.status_code == 409 and r.json()["code"] == "no_documents"


# ---------------------------------------------------------------- ask
def test_ask_returns_grounded_answer_with_sources(client, uploaded, mock_llm):
    r = client.post("/api/ask", json={"question": "What is FAISS used for?"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["answer"].startswith("Mock answer") and body["model"] == "mock-model"
    src = body["sources"][0]
    assert set(src) >= {"document_name", "page", "chunk_id", "score", "passage", "document_id"}
    assert src["document_name"] == "notes.txt" and src["page"] is None and src["score"] > 0
    question_sent, sources_sent = mock_llm[0]
    assert "FAISS" in question_sent and sources_sent


def test_ask_accepts_camel_case_request(client, uploaded, mock_llm):
    r = client.post("/api/ask", json={"question": "What is Chroma?", "topK": 2, "documentIds": [uploaded["id"]]})
    assert r.status_code == 200 and len(r.json()["sources"]) <= 2


def test_ask_empty_question(client, uploaded):
    for q in ("", "   "):
        r = client.post("/api/ask", json={"question": q})
        assert r.status_code == 400 and r.json()["code"] == "empty_question"


def test_ask_without_documents(client):
    r = client.post("/api/ask", json={"question": "hello?"})
    assert r.status_code == 409 and r.json()["code"] == "no_documents"


def test_ask_missing_api_key(client, uploaded, monkeypatch):
    from app.config import get_settings
    from dataclasses import replace

    monkeypatch.setattr(llm, "get_settings", lambda: replace(get_settings(), groq_api_key=""))
    r = client.post("/api/ask", json={"question": "What is FAISS?"})
    assert r.status_code == 503 and r.json()["code"] == "missing_api_key"
    assert client.get("/api/history").json()["total"] == 0  # failures are not saved as answers


def test_ask_provider_failure(client, uploaded, monkeypatch):
    def boom(question, sources):
        raise llm.LLMError(502, "provider_error", "Groq returned HTTP 500")

    monkeypatch.setattr(llm, "generate_answer", boom)
    r = client.post("/api/ask", json={"question": "What is FAISS?"})
    assert r.status_code == 502 and r.json()["code"] == "provider_error"


def test_groq_http_layer_with_fake_transport(client, uploaded, monkeypatch):
    import httpx

    def fake_post(url, headers, payload, timeout):
        assert headers["Authorization"] == "Bearer test-key-not-real"
        assert payload["messages"][1]["content"].count("Source: notes.txt") >= 1
        return httpx.Response(200, json={"model": "llama-test", "choices": [{"message": {"content": "FAISS does search [1]."}}]})

    monkeypatch.setattr(llm, "_post", fake_post)
    r = client.post("/api/ask", json={"question": "What is FAISS used for?"})
    assert r.status_code == 200 and r.json()["answer"] == "FAISS does search [1]." and r.json()["model"] == "llama-test"


@pytest.mark.parametrize("status,code,http", [(401, "provider_auth_error", 502), (429, "provider_rate_limited", 429), (500, "provider_error", 502)])
def test_groq_error_mapping(client, uploaded, monkeypatch, status, code, http):
    import httpx

    monkeypatch.setattr(llm, "_post", lambda *a, **k: httpx.Response(status, json={"error": {"message": "nope"}}))
    r = client.post("/api/ask", json={"question": "What is FAISS used for?"})
    assert r.status_code == http and r.json()["code"] == code


# ---------------------------------------------------------------- history & export
@pytest.fixture()
def two_answers(client, uploaded, mock_llm):
    ids = []
    for q in ("What is FAISS?", "=SUM(1,1) injection test?"):
        ids.append(client.post("/api/ask", json={"question": q}).json()["id"])
    return ids


def test_history_list_get_delete(client, two_answers):
    h = client.get("/api/history").json()
    assert h["total"] == 2 and h["history"][0]["id"] == two_answers[1]  # newest first
    one = client.get(f"/api/history/{two_answers[0]}")
    assert one.status_code == 200 and one.json()["question"] == "What is FAISS?"
    assert client.get("/api/history/9999").status_code == 404
    assert client.delete("/api/history").json()["deleted"] == 2
    assert client.get("/api/history").json()["total"] == 0


def test_export_all_formats(client, two_answers):
    j = client.get("/api/history/export", params={"format": "json"})
    assert j.status_code == 200 and len(json.loads(j.content)) == 2
    assert "attachment" in j.headers["content-disposition"]

    c = client.get("/api/history/export", params={"format": "csv"})
    assert c.status_code == 200 and b"What is FAISS?" in c.content and b"'=SUM" in c.content

    x = client.get("/api/history/export", params={"format": "xlsx"})
    assert x.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(x.content))["History"]
    assert ws.max_row == 3 and ws["A1"].value == "id"


def test_export_single_record(client, two_answers):
    rid = two_answers[0]
    j = client.get(f"/api/history/{rid}/export", params={"format": "json"})
    assert json.loads(j.content)["id"] == rid
    assert client.get(f"/api/history/{rid}/export", params={"format": "csv"}).status_code == 200
    assert client.get(f"/api/history/{rid}/export", params={"format": "xlsx"}).status_code == 200
    assert client.get("/api/history/9999/export").status_code == 404


def test_export_invalid_format(client, two_answers):
    r = client.get("/api/history/export", params={"format": "pdf"})
    assert r.status_code == 400 and r.json()["code"] == "invalid_format"


def test_export_empty_history(client):
    for fmt in ("json", "csv", "xlsx"):
        assert client.get("/api/history/export", params={"format": fmt}).status_code == 200


# ---------------------------------------------------------------- experiments
def test_run_and_list_experiments(client, uploaded):
    r = client.post(
        "/api/experiments/run",
        json={"chunk_size": 200, "chunk_overlap": 40, "vector_store": "FAISS", "top_k": 2,
              "queries": ["What is FAISS?", "Where is history stored?"], "notes": "first run"},
    )
    assert r.status_code == 201, r.text
    exp = r.json()
    assert exp["config"]["chunk_size"] == 200 and exp["notes"] == "first run"
    assert exp["summary"]["query_count"] == 2 and exp["summary"]["total_chunks"] > 0
    assert exp["summary"]["effective_vector_store"] in ("faiss", "numpy")
    assert exp["summary"]["semantic"] is False  # hashing backend in tests; must be reported honestly
    for res in exp["results"]:
        assert len(res["hits"]) <= 2 and res["latency_ms"] >= 0
    assert not {"precision", "recall", "accuracy", "relevance", "ndcg"} & set(exp["summary"])  # nothing fabricated
    listing = client.get("/api/experiments").json()
    assert listing["total"] == 1 and listing["experiments"][0]["id"] == exp["id"]


def test_experiment_camel_case_and_invalid_overlap(client, uploaded):
    ok = client.post("/api/experiments/run", json={"chunkSize": 300, "chunkOverlap": 50, "vectorStore": "numpy", "topK": 3, "query": "chunk overlap"})
    assert ok.status_code == 201
    bad = client.post("/api/experiments/run", json={"chunk_size": 100, "chunk_overlap": 100, "queries": ["x"]})
    assert bad.status_code == 400 and bad.json()["code"] == "invalid_experiment_config"


def test_experiment_without_documents(client):
    r = client.post("/api/experiments/run", json={"queries": ["x"]})
    assert r.status_code == 409 and r.json()["code"] == "no_documents"


def test_experiment_does_not_change_live_index(client, uploaded, mock_llm):
    before = client.get("/api/documents").json()["documents"][0]["chunk_count"]
    client.post("/api/experiments/run", json={"chunk_size": 100, "chunk_overlap": 10, "queries": ["FAISS"]})
    assert client.get("/api/documents").json()["documents"][0]["chunk_count"] == before
    assert client.post("/api/ask", json={"question": "What is FAISS?"}).status_code == 200


def test_validation_error_shape(client):
    r = client.post("/api/ask", json={})
    assert r.status_code == 422 and r.json()["code"] == "validation_error" and isinstance(r.json()["detail"], str)
